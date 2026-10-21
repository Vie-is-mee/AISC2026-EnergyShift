"""Strict local CSV/XLSX -> factory-1/factory-2 review adapter; never optimizes during import.

A workbook uses named sheets; a single CSV carries a `section` column.
All required machine, tariff, power and baseline facts must be supplied.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .factory_core import _validate, _package
from .factory_routing import RoutingInput, _routing_issues, _to_factory1
from pydantic import ValidationError

MAX_UPLOAD_BYTES = 350_000
MAX_ROWS = 1600
SHEETS = {
    "Settings": ("key", "value"),
    "Machines": ("machine_id", "available_start", "available_end"),
    "Orders": ("work_order_id", "release", "due"),
    "Tasks": ("task_id", "work_order_id", "release", "due", "setup_family"),
    "Capabilities": ("task_id", "machine_id", "duration_minutes", "average_kw",
                     "power_source", "reference"),
    "Schedule": ("task_id", "machine_id", "start"),
    "Tariffs": ("start", "end", "band", "vnd_per_kwh"),
    "Precedence": ("before", "after", "min_lag_minutes", "work_order_id"),
    "Resources": ("resource_id", "capacity", "available_start", "available_end"),
    "Demands": ("task_id", "resource_id", "amount"),
    "SetupPolicy": ("machine_id", "initial_family"),
    "Transitions": ("machine_id", "from_family", "to_family", "duration_minutes",
                    "flat_cost_vnd", "reference"),
    "SetupDemands": ("machine_id", "from_family", "to_family", "resource_id", "amount"),
}
COLUMNS = ("section",) + tuple(dict.fromkeys(k for names in SHEETS.values() for k in names))
REQUIRED_SETTINGS = ("schema_version", "plant_alias", "industry", "source_type",
                     "timezone", "currency", "horizon_start", "horizon_end",
                     "step_minutes", "tariff_reference", "setup_assumption",
                     "unsupported_constraints")
OPTIONAL_SETTINGS = ("buffer_assumption",)


def _error(code, details=None, *, draft=None, manifest=None):
    result = {"status": "BLOCKED", "code": code, "issues": details or [],
              "ready_for_analysis": False}
    if draft is not None:
        result["draft"] = draft
    if manifest is not None:
        result["manifest"] = manifest
    return result


def _value(v):
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        # Excel timestamps lacking +07:00 must not be silently assigned a timezone.
        return v.isoformat()
    return str(v).strip()


def _nonblank(row, key, issues, section, index):
    s = row.get(key, "")
    if not s:
        issues.append({"code": "IMPORT_REQUIRED_CELL", "section": section, "row": index, "field": key})
    return s


def _number(row, field, issues, section, index, *, integer=True):
    val = _nonblank(row, field, issues, section, index)
    try:
        if integer:
            if not val or not val.lstrip("+-").isdigit():
                raise ValueError("not integer")
            return int(val)
        if not val or "," in val:
            raise ValueError("not number")
        number = float(val)
        if not (-1e15 < number < 1e15):
            raise ValueError("non finite / out of bounds")
        return number
    except (ValueError, OverflowError):
        issues.append({"code": "IMPORT_BAD_NUMBER", "section": section, "row": index, "field": field})
        return 0


def _read_csv(raw):
    try:
        text = raw.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError:
        raise ValueError("IMPORT_UTF8_REQUIRED")
    reader = csv.DictReader(io.StringIO(text, newline=""), strict=True)
    if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
        raise ValueError("IMPORT_HEADER_INVALID")
    if set(reader.fieldnames) != set(COLUMNS):
        raise ValueError("IMPORT_HEADER_INVALID")
    groups = {sheet: [] for sheet in SHEETS}
    for index, line in enumerate(reader, 2):
        if index > MAX_ROWS + 1:
            raise ValueError("IMPORT_ROW_LIMIT")
        if None in line or any(value is None for value in line.values()):
            raise ValueError("IMPORT_ROW_SHAPE_INVALID")
        name = _value(line.get("section"))
        if name not in groups:
            raise ValueError("IMPORT_UNKNOWN_SECTION")
        unused = [key for key in COLUMNS if key != "section" and key not in SHEETS[name] and _value(line[key])]
        if unused:
            raise ValueError("IMPORT_UNUSED_COLUMN_DATA")
        groups[name].append(({key: _value(line[key]) for key in SHEETS[name]}, index))
    return groups


def _read_xlsx(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist()) > 120 or sum(i.file_size for i in z.infolist()) > 5_000_000:
                raise ValueError("IMPORT_XLSX_EXPANDED_LIMIT")
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except Exception as e:
        if isinstance(e, ValueError) and str(e) == "IMPORT_XLSX_EXPANDED_LIMIT":
            raise
        # Reject damaged archives / invalid spreadsheet XML without exposing internals.
        raise ValueError("IMPORT_XLSX_INVALID") from e
    try:
        if not set(SHEETS).issuperset(book.sheetnames):
            raise ValueError("IMPORT_UNKNOWN_SECTION")
        groups = {name: [] for name in SHEETS}
        total = 0
        for name in book.sheetnames:
            sheet = book[name]
            rows = sheet.iter_rows(values_only=False)
            raw_header = next(rows, None)
            header = tuple(cell.value for cell in raw_header) if raw_header is not None else None
            if header is None or tuple(map(_value, header)) != SHEETS[name]:
                raise ValueError("IMPORT_HEADER_INVALID")
            for i, cells in enumerate(rows, 2):
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("IMPORT_FORMULAS_UNSUPPORTED")
                values = tuple(cell.value for cell in cells)
                if all(v is None or _value(v) == "" for v in values):
                    continue
                total += 1
                if total > MAX_ROWS:
                    raise ValueError("IMPORT_ROW_LIMIT")
                if len(values) != len(SHEETS[name]):
                    raise ValueError("IMPORT_ROW_SHAPE_INVALID")
                groups[name].append((dict(zip(SHEETS[name], map(_value, values))), i))
        return groups
    finally:
        book.close()


def _build(groups):
    issues = []
    settings = {}
    for row, n in groups["Settings"]:
        key = _nonblank(row, "key", issues, "Settings", n)
        if key in settings:
            issues.append({"code": "IMPORT_DUPLICATE_SETTING", "field": key})
        if key not in REQUIRED_SETTINGS + OPTIONAL_SETTINGS:
            issues.append({"code": "IMPORT_UNKNOWN_SETTING", "field": key})
        settings[key] = row["value"]
    for key in REQUIRED_SETTINGS:
        if not settings.get(key) and key != "unsupported_constraints":
            issues.append({"code": "IMPORT_MISSING_SETTING", "field": key})
    version = settings.get("schema_version")
    if version not in ("factory-1", "factory-2"):
        issues.append({"code": "IMPORT_SCHEMA_UNSUPPORTED", "field": "schema_version"})
    if version == "factory-2" and not settings.get("buffer_assumption"):
        issues.append({"code": "IMPORT_MISSING_SETTING", "field": "buffer_assumption"})
    if settings.get("unsupported_constraints"):
        # Explicit nonempty unsupported constraints cannot be silently dropped.
        settings["unsupported_constraints"] = [s.strip() for s in settings["unsupported_constraints"].split(";")]
    else:
        settings["unsupported_constraints"] = []
    steps = {"step_minutes": _number({"step_minutes": settings.get("step_minutes", "")},
                                       "step_minutes", issues, "Settings", 0)}
    data = {
        "schema_version": version, "metadata": {
            "plant_alias": settings.get("plant_alias", ""),
            "industry": settings.get("industry", ""),
            "source_type": settings.get("source_type", ""),
            "timezone": settings.get("timezone", ""), "currency": settings.get("currency", "")
        },
        "horizon": {"start": settings.get("horizon_start", ""), "end": settings.get("horizon_end", "")},
        "step_minutes": steps["step_minutes"], "machines": [], "tariffs": [],
        "tariff_reference": settings.get("tariff_reference", ""),
        "setup_assumption": settings.get("setup_assumption", ""),
        "unsupported_constraints": settings["unsupported_constraints"],
        "resources": [], "current_schedule": [],
    }
    if version == "factory-2":
        data["buffer_assumption"] = settings.get("buffer_assumption", "")
        data["work_orders"] = []
    else:
        data["jobs"] = []
        data["precedence"] = []

    def indexed(section, keys):
        result = {}
        for row, i in groups[section]:
            key = tuple(row[k] for k in keys)
            if not all(key):
                for k in keys:
                    _nonblank(row, k, issues, section, i)
                continue
            if key in result:
                issues.append({"code": "IMPORT_DUPLICATE_KEY", "section": section, "row": i})
                continue
            result[key] = (row, i)
        return result

    machines = {}
    for row, i in groups["Machines"]:
        mid = _nonblank(row, "machine_id", issues, "Machines", i)
        window = {"start": _nonblank(row, "available_start", issues, "Machines", i),
                  "end": _nonblank(row, "available_end", issues, "Machines", i)}
        machine = machines.setdefault(mid, {"id": mid, "available": []})
        machine["available"].append(window)
    policy = indexed("SetupPolicy", ("machine_id",))
    for (mid,), (row, i) in policy.items():
        if mid not in machines:
            issues.append({"code": "IMPORT_UNKNOWN_MACHINE", "section": "SetupPolicy", "row": i})
        else:
            machines[mid]["changeovers"] = {
                "initial_family": _nonblank(row, "initial_family", issues, "SetupPolicy", i),
                "transitions": []}
    demands = defaultdict(dict)
    for (task, rid), (row, i) in indexed("Demands", ("task_id", "resource_id")).items():
        demands[task][rid] = _number(row, "amount", issues, "Demands", i)
    setup_demands = defaultdict(dict)
    for (mid, frm, to, rid), (row, i) in indexed(
            "SetupDemands", ("machine_id", "from_family", "to_family", "resource_id")).items():
        setup_demands[(mid, frm, to)][rid] = _number(row, "amount", issues, "SetupDemands", i)
    transition_rows = indexed("Transitions", ("machine_id", "from_family", "to_family"))
    for (mid, frm, to), (row, i) in transition_rows.items():
        if mid not in machines or "changeovers" not in machines[mid]:
            issues.append({"code": "IMPORT_MISSING_SETUP_POLICY", "section": "Transitions", "row": i})
            continue
        machines[mid]["changeovers"]["transitions"].append({
            "from_family": frm, "to_family": to,
            "duration_minutes": _number(row, "duration_minutes", issues, "Transitions", i),
            "flat_cost_vnd": _number(row, "flat_cost_vnd", issues, "Transitions", i),
            "reference": _nonblank(row, "reference", issues, "Transitions", i),
            "demands": setup_demands[(mid, frm, to)]})
    for mid, frm, to in setup_demands:
        if (mid, frm, to) not in transition_rows:
            issues.append({"code": "IMPORT_SETUP_DEMAND_UNKNOWN_TRANSITION", "machine_id": mid})
    data["machines"] = list(machines.values())

    capabilities = defaultdict(dict)
    for (task, mid), (row, i) in indexed("Capabilities", ("task_id", "machine_id")).items():
        capabilities[task][mid] = {
            "duration_minutes": _number(row, "duration_minutes", issues, "Capabilities", i),
            "average_kw": _number(row, "average_kw", issues, "Capabilities", i, integer=False),
            "power_source": _nonblank(row, "power_source", issues, "Capabilities", i),
            "reference": _nonblank(row, "reference", issues, "Capabilities", i)}
    task_map = {}
    for (tid,), (row, i) in indexed("Tasks", ("task_id",)).items():
        task = {"id": tid, "eligible": capabilities[tid], "demands": demands[tid]}
        if row["setup_family"]:
            task["setup_family"] = row["setup_family"]
        if version == "factory-2":
            _nonblank(row, "work_order_id", issues, "Tasks", i)
        else:
            if row["work_order_id"]:
                issues.append({"code": "IMPORT_UNEXPECTED_WORK_ORDER", "row": i})
            task["release"] = _nonblank(row, "release", issues, "Tasks", i)
            task["due"] = _nonblank(row, "due", issues, "Tasks", i)
        task_map[tid] = (task, row)
    for task in capabilities:
        if task not in task_map:
            issues.append({"code": "IMPORT_CAPABILITY_UNKNOWN_TASK", "task_id": task})
    for task in demands:
        if task not in task_map:
            issues.append({"code": "IMPORT_DEMAND_UNKNOWN_TASK", "task_id": task})
    order_map = {}
    for (oid,), (row, i) in indexed("Orders", ("work_order_id",)).items():
        order_map[oid] = {
            "id": oid, "release": _nonblank(row, "release", issues, "Orders", i),
            "due": _nonblank(row, "due", issues, "Orders", i),
            "operations": [], "precedence": []}
    if version == "factory-2":
        for tid, (task, row) in task_map.items():
            oid = row["work_order_id"]
            if oid not in order_map:
                issues.append({"code": "IMPORT_UNKNOWN_WORK_ORDER", "task_id": tid})
                continue
            order_map[oid]["operations"].append(task)
        data["work_orders"] = list(order_map.values())
    else:
        if order_map:
            issues.append({"code": "IMPORT_UNEXPECTED_ORDERS"})
        data["jobs"] = [item[0] for item in task_map.values()]

    for (before, after), (row, i) in indexed("Precedence", ("before", "after")).items():
        edge = {"before": before, "after": after,
                "min_lag_minutes": _number(row, "min_lag_minutes", issues, "Precedence", i)}
        if version == "factory-2":
            oid = _nonblank(row, "work_order_id", issues, "Precedence", i)
            if oid not in order_map:
                issues.append({"code": "IMPORT_UNKNOWN_WORK_ORDER", "section": "Precedence", "row": i})
            else:
                order_map[oid]["precedence"].append(edge)
        else:
            if row["work_order_id"]:
                issues.append({"code": "IMPORT_UNEXPECTED_WORK_ORDER", "row": i})
            data["precedence"].append(edge)

    resources = {}
    for row, i in groups["Resources"]:
        rid = _nonblank(row, "resource_id", issues, "Resources", i)
        cap = _number(row, "capacity", issues, "Resources", i)
        if rid in resources and resources[rid]["capacity"] != cap:
            issues.append({"code": "IMPORT_RESOURCE_CAPACITY_CONFLICT", "row": i})
        pool = resources.setdefault(rid, {"id": rid, "capacity": cap, "available": []})
        pool["available"].append({
            "start": _nonblank(row, "available_start", issues, "Resources", i),
            "end": _nonblank(row, "available_end", issues, "Resources", i)})
    data["resources"] = list(resources.values())
    for (tid,), (row, i) in indexed("Schedule", ("task_id",)).items():
        data["current_schedule"].append({
            "operation_id" if version == "factory-2" else "job_id": tid,
            "machine_id": _nonblank(row, "machine_id", issues, "Schedule", i),
            "start": _nonblank(row, "start", issues, "Schedule", i)})
    for row, i in groups["Tariffs"]:
        data["tariffs"].append({
            "start": _nonblank(row, "start", issues, "Tariffs", i),
            "end": _nonblank(row, "end", issues, "Tariffs", i),
            "band": _nonblank(row, "band", issues, "Tariffs", i),
            "vnd_per_kwh": _number(row, "vnd_per_kwh", issues, "Tariffs", i, integer=False)})
    return data, issues


def _preflight(draft):
    if draft["schema_version"] == "factory-1":
        data, blocked = _validate(draft)
        if blocked:
            return blocked["issues"] or [{"code": blocked["code"]}]
        baseline = [{"job_id": a.job_id, "machine_id": a.machine_id, "start": a.start}
                    for a in data.current_schedule]
        proof = _package(data, baseline)
        return proof["verification"]["issues"] if proof["status"] != "PASS" else []
    try:
        from .factory_routing import RoutingInput
        data = RoutingInput.model_validate(draft)
    except ValidationError as e:
        return [{"code": "SCHEMA_INVALID", "field": ".".join(map(str, x["loc"]))}
                for x in e.errors(include_input=False, include_url=False)][:30]
    issues, _ = _routing_issues(data)
    if issues:
        return issues
    flat = _to_factory1(data)
    return _preflight(flat)


def inspect_spreadsheet(raw, file_type):
    if not isinstance(raw, bytes) or len(raw) > MAX_UPLOAD_BYTES or not raw:
        return _error("IMPORT_FILE_SIZE", [{"code": "FILE_SIZE_LIMIT", "bytes": len(raw) if raw else 0}])
    if file_type not in ("csv", "xlsx"):
        return _error("IMPORT_FILE_TYPE")
    try:
        groups = _read_csv(raw) if file_type == "csv" else _read_xlsx(raw)
        draft, problems = _build(groups)
    except (ValueError, csv.Error) as error:
        return _error(str(error) if str(error).startswith("IMPORT_") else "IMPORT_INVALID")
    source_hash = "sha256:" + hashlib.sha256(raw).hexdigest()
    manifest = {"sha256": source_hash, "file_type": file_type,
                "sections": {key: len(value) for key, value in groups.items()},
                "source_type": draft["metadata"]["source_type"],
                "schema_version": draft["schema_version"],
                "provenance": "USER_UPLOADED_NOT_VERIFIED"}
    if problems:
        return _error("IMPORT_MAPPING_INVALID", problems[:60], draft=draft, manifest=manifest)
    constraints = _preflight(draft)
    if constraints:
        return _error("IMPORT_PREFLIGHT_BLOCKED", constraints[:60], draft=draft, manifest=manifest)
    return {"status": "READY_FOR_REVIEW", "ready_for_analysis": True,
            "draft": draft, "manifest": manifest,
            "issues": [], "review_required": True,
            "message": "Review source facts, operation routes, tariffs and assumptions before explicit acceptance."}


def to_sections(factory):
    """Create a readable, editable sample template from a declared synthetic input."""
    base = {name: [] for name in SHEETS}
    version = factory["schema_version"]
    for key, value in [
        ("schema_version", version), ("plant_alias", factory["metadata"]["plant_alias"]),
        ("industry", factory["metadata"]["industry"]), ("source_type", factory["metadata"]["source_type"]),
        ("timezone", factory["metadata"]["timezone"]), ("currency", factory["metadata"]["currency"]),
        ("horizon_start", factory["horizon"]["start"]), ("horizon_end", factory["horizon"]["end"]),
        ("step_minutes", factory["step_minutes"]), ("tariff_reference", factory["tariff_reference"]),
        ("setup_assumption", factory["setup_assumption"]),
        ("unsupported_constraints", ";".join(factory.get("unsupported_constraints", []))),
    ]:
        base["Settings"].append({"key": key, "value": str(value)})
    if version == "factory-2":
        base["Settings"].append({"key": "buffer_assumption", "value": factory["buffer_assumption"]})

    for machine in factory["machines"]:
        for w in machine["available"]:
            base["Machines"].append({"machine_id": machine["id"],
                                     "available_start": w["start"], "available_end": w["end"]})
        policy = machine.get("changeovers")
        if policy:
            base["SetupPolicy"].append({"machine_id": machine["id"],
                                        "initial_family": policy["initial_family"]})
            for t in policy["transitions"]:
                base["Transitions"].append({"machine_id": machine["id"], "from_family": t["from_family"],
                    "to_family": t["to_family"], "duration_minutes": t["duration_minutes"],
                    "flat_cost_vnd": t["flat_cost_vnd"], "reference": t["reference"]})
                for rid, amount in t.get("demands", {}).items():
                    base["SetupDemands"].append({"machine_id": machine["id"], "from_family": t["from_family"],
                        "to_family": t["to_family"], "resource_id": rid, "amount": amount})
    if version == "factory-1":
        tasks = [(j, None) for j in factory["jobs"]]
        edges = [(p, None) for p in factory.get("precedence", [])]
        assignment_key = "job_id"
    else:
        tasks, edges = [], []
        for o in factory["work_orders"]:
            base["Orders"].append({"work_order_id": o["id"], "release": o["release"], "due": o["due"]})
            tasks.extend((p, o["id"]) for p in o["operations"])
            edges.extend((e, o["id"]) for e in o["precedence"])
        assignment_key = "operation_id"
    for task, oid in tasks:
        base["Tasks"].append({"task_id": task["id"], "work_order_id": oid or "",
                              "release": task.get("release", ""), "due": task.get("due", ""),
                              "setup_family": task.get("setup_family", "") or ""})
        for mid, p in task["eligible"].items():
            base["Capabilities"].append({"task_id": task["id"], "machine_id": mid,
                "duration_minutes": p["duration_minutes"], "average_kw": p["average_kw"],
                "power_source": p["power_source"], "reference": p["reference"]})
        for rid, amount in task.get("demands", {}).items():
            base["Demands"].append({"task_id": task["id"], "resource_id": rid, "amount": amount})
    for edge, oid in edges:
        base["Precedence"].append({"before": edge["before"], "after": edge["after"],
                                   "min_lag_minutes": edge.get("min_lag_minutes", 0),
                                   "work_order_id": oid or ""})
    for s in factory["current_schedule"]:
        base["Schedule"].append({"task_id": s[assignment_key],
                                 "machine_id": s["machine_id"], "start": s["start"]})
    for tariff in factory["tariffs"]:
        base["Tariffs"].append(dict(tariff))
    for r in factory.get("resources", []):
        for w in r["available"]:
            base["Resources"].append({"resource_id": r["id"], "capacity": r["capacity"],
                                      "available_start": w["start"], "available_end": w["end"]})
    return base


def template_bytes(factory, file_type):
    tables = to_sections(factory)
    if file_type == "csv":
        output = io.StringIO(newline="")
        writer = csv.DictWriter(output, fieldnames=COLUMNS, lineterminator="\r\n")
        writer.writeheader()
        for name, rows in tables.items():
            for row in rows:
                writer.writerow({"section": name, **row})
        return ("\ufeff" + output.getvalue()).encode("utf-8")
    if file_type != "xlsx":
        raise ValueError("IMPORT_FILE_TYPE")
    book = Workbook()
    book.remove(book.active)
    for name, names in SHEETS.items():
        sheet = book.create_sheet(name)
        sheet.append(names)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(names))}{max(2,len(tables[name])+1)}"
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="17365D")
            c.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 28
        for i, col in enumerate(names, 1):
            sheet.column_dimensions[get_column_letter(i)].width = min(32, max(13, len(col) + 3))
        for row in tables[name]:
            sheet.append([row.get(col, "") for col in names])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()
